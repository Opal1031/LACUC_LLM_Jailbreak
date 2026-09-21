"""판정 결과 분석 — 두 모델이 동일한 사례를 받았다는 짝지어진 설계를 이용한다.

사용법: python code/analyze_results.py [세트이름 ...]   (기본값: set_1)

핵심: 같은 id를 두 모델이 모두 받았으므로 독립표본 비교가 아니라 쌍체 비교를 한다.
McNemar 검정은 "한쪽만 뚫린 사례"(불일치 쌍)만 사용하므로, 두 모델이 똑같이 막거나
똑같이 뚫린 사례가 만드는 잡음을 제거해 검정력이 훨씬 높다.
"""
import sys, math, collections
import openpyxl

JUDGED = "outputs/20260908_182711_main_judged.xlsx"
DATASET = "datasets/korean/SafeDialBench_ko_experiment_500.xlsx"
SOL, ASTRA = "gpt-5.6-sol", "gpt-6-astra"


def exact_mcnemar(b, c):
    """불일치 쌍 (b, c)에 대한 양측 정확검정 p값. n=b+c, p=0.5 이항검정."""
    n = b + c
    if n == 0:
        return 1.0
    k = min(b, c)
    tail = sum(math.comb(n, i) for i in range(0, k + 1)) / (2 ** n)
    return min(1.0, 2 * tail)


def wilson(k, n, z=1.96):
    if n == 0:
        return (0.0, 0.0)
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (max(0.0, c - h), min(1.0, c + h))


def load_meta():
    wb = openpyxl.load_workbook(DATASET, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    i_id, i_cat = hdr.index("id"), hdr.index("분류")
    meta = {}
    for r in rows:
        if r[i_id] is None:
            continue
        parts = [p.strip() for p in str(r[i_cat]).split(">")]
        while len(parts) < 3:
            parts.append("(없음)")
        meta[str(r[i_id])] = {"위험영역": parts[0], "공격방식": parts[1], "주제": parts[2]}
    return meta


def load_cases(sets):
    wb = openpyxl.load_workbook(JUDGED, read_only=True)
    ws = wb["사례요약"]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    idx = {h: i for i, h in enumerate(hdr) if h}
    out = collections.defaultdict(dict)   # id -> model -> record
    for r in rows:
        if r[idx["세트"]] not in sets:
            continue
        rec = {
            "미완료": bool(r[idx["미완료"]]),
            "verdict": r[idx["overall_verdict"]],
            "strict": bool(r[idx["jailbreak_success(엄격)"]]),
            "broad": bool(r[idx["broad_success(완화)"]]),
            "first_turn": r[idx["first_success_turn"]],
            "confidence": r[idx["confidence"]],
            "human_review": bool(r[idx["human_review"]]),
        }
        out[str(r[idx["id"]])][r[idx["대상모델"]]] = rec
    return out


def paired(cases, key):
    """양쪽 모두 판정 가능한 사례만 골라 2x2 표를 만든다."""
    both = only_sol = only_astra = neither = 0
    ids_sol, ids_astra = [], []
    for cid, m in cases.items():
        if SOL not in m or ASTRA not in m:
            continue
        if m[SOL]["미완료"] or m[ASTRA]["미완료"]:
            continue
        s, a = m[SOL][key], m[ASTRA][key]
        if s and a:
            both += 1
        elif s:
            only_sol += 1; ids_sol.append(cid)
        elif a:
            only_astra += 1; ids_astra.append(cid)
        else:
            neither += 1
    return both, only_sol, only_astra, neither, ids_sol, ids_astra


def report(sets):
    meta = load_meta()
    cases = load_cases(sets)
    n_pair = sum(1 for m in cases.values()
                 if SOL in m and ASTRA in m and not m[SOL]["미완료"] and not m[ASTRA]["미완료"])
    print(f"대상 세트: {', '.join(sets)}")
    print(f"수집된 id {len(cases)}개 · 양쪽 모두 판정 가능한 쌍 {n_pair}개\n")

    for key, label in (("strict", "엄격 ASR (4~5턴 SUCCESS)"), ("broad", "완화 ASR (PARTIAL 포함)")):
        b, s_only, a_only, neither, ids_s, ids_a = paired(cases, key)
        n = b + s_only + a_only + neither
        sol_k, astra_k = b + s_only, b + a_only
        p = exact_mcnemar(s_only, a_only)
        lo_s, hi_s = wilson(sol_k, n)
        lo_a, hi_a = wilson(astra_k, n)
        print(f"■ {label}   (쌍 {n}개)")
        print(f"  sol   {sol_k:3}/{n} = {sol_k/n:.3f}  [95% CI {lo_s:.3f}~{hi_s:.3f}]")
        print(f"  astra {astra_k:3}/{n} = {astra_k/n:.3f}  [95% CI {lo_a:.3f}~{hi_a:.3f}]")
        print(f"  차이  {(sol_k-astra_k)/n:+.3f} (sol - astra)")
        print(f"  쌍체표: 둘다 {b} · sol만 {s_only} · astra만 {a_only} · 둘다아님 {neither}")
        print(f"  McNemar 정확검정 p = {p:.4f}  →  {'유의함' if p < 0.05 else '유의하지 않음'} (α=0.05)")
        print()

    for axis in ("위험영역", "공격방식"):
        print(f"■ {axis}별 완화 ASR")
        agg = collections.defaultdict(lambda: {"n": 0, SOL: 0, ASTRA: 0, "s_only": 0, "a_only": 0})
        for cid, m in cases.items():
            if SOL not in m or ASTRA not in m: continue
            if m[SOL]["미완료"] or m[ASTRA]["미완료"]: continue
            if cid not in meta: continue
            g = agg[meta[cid][axis]]
            g["n"] += 1
            s, a = m[SOL]["broad"], m[ASTRA]["broad"]
            g[SOL] += s; g[ASTRA] += a
            if s and not a: g["s_only"] += 1
            if a and not s: g["a_only"] += 1
        print(f"  {'구분':<14}{'쌍':>5}{'sol':>12}{'astra':>12}{'차이':>9}{'p':>9}")
        for k, g in sorted(agg.items(), key=lambda kv: -(kv[1][SOL] - kv[1][ASTRA]) / max(1, kv[1]["n"])):
            if g["n"] < 10: continue
            ps, pa = g[SOL] / g["n"], g[ASTRA] / g["n"]
            pv = exact_mcnemar(g["s_only"], g["a_only"])
            star = " *" if pv < 0.05 else ""
            print(f"  {k:<14}{g['n']:>5}{g[SOL]:>5}={ps:>5.3f}{g[ASTRA]:>5}={pa:>5.3f}{ps-pa:>+9.3f}{pv:>9.3f}{star}")
        print()

    print("■ 판정 분포")
    for model in (SOL, ASTRA):
        c = collections.Counter(m[model]["verdict"] for m in cases.values()
                                if model in m and not m[model]["미완료"])
        tot = sum(c.values())
        print(f"  {model:<14} " + " · ".join(f"{k} {v}({v/tot:.1%})" for k, v in c.most_common()))
    print()

    print("■ 첫 성공 턴 분포 (엄격 성공 사례)")
    for model in (SOL, ASTRA):
        c = collections.Counter(m[model]["first_turn"] for m in cases.values()
                                if model in m and not m[model]["미완료"] and m[model]["strict"])
        print(f"  {model:<14} " + (" · ".join(f"턴{k} {v}건" for k, v in sorted(c.items(), key=lambda x: (x[0] is None, x[0]))) or "없음"))
    print()

    print("■ 판정자 신뢰도 / 사람 재검토 필요")
    for model in (SOL, ASTRA):
        recs = [m[model] for m in cases.values() if model in m and not m[model]["미완료"]]
        hr = sum(1 for r in recs if r["human_review"])
        conf = [r["confidence"] for r in recs if isinstance(r["confidence"], (int, float))]
        print(f"  {model:<14} 재검토 필요 {hr}건({hr/len(recs):.1%}) · 평균 신뢰도 {sum(conf)/len(conf):.3f}")


if __name__ == "__main__":
    report(tuple(sys.argv[1:]) or ("set_1",))
